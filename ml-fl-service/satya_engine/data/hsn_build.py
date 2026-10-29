"""Build only from explicit local source files; never infer tax rates."""
import csv
import gzip
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
from zipfile import ZipFile

from ..india.hsn_master import HsnArtifact

MAX_SOURCE_BYTES = 64 * 1024 * 1024
MAX_ARTIFACT_BYTES = 5 * 1024 * 1024


def source_metadata(path, name, url, retrieved_at):
    host = urlparse(url).hostname or ""
    if not url.startswith("https://") or not host.endswith((".gov.in", ".nic.in")):
        raise ValueError("Reference source URL must be a primary government HTTPS source")
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(64 * 1024):
            digest.update(chunk)
    return dict(name=name, url=url, retrievedAt=retrieved_at, sha256=digest.hexdigest())


def directory_rows(path):
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Source exceeds 64 MiB limit; supply a smaller relevant subset")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            yield from csv.reader(stream)
    elif path.suffix.lower() == ".xlsx":
        # Check expanded size before openpyxl reads any ZIP contents.
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_SOURCE_BYTES:
                raise ValueError("Expanded workbook exceeds 64 MiB limit")
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            yield from workbook.active.iter_rows(values_only=True)
        finally:
            workbook.close()
    else:
        raise ValueError("Use an explicit CSV or XLSX directory file")


def reviewed_rates(path):
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Rate extraction exceeds source size cap")
    data = json.loads(path.read_text())
    if not data.get("sources") or data.get("extractionReviewed") is not True:
        raise ValueError("Rate rules require primary source files and extractionReviewed=true")
    sources = []
    for entry in data["sources"]:
        local_path = path.parent / entry["localPath"]
        if local_path.stat().st_size > MAX_SOURCE_BYTES:
            raise ValueError("Primary rate source exceeds source size cap")
        metadata = source_metadata(local_path, entry["name"], entry["url"], entry["retrievedAt"])
        if metadata["sha256"] != entry["sha256"]:
            raise ValueError("Primary source checksum does not match reviewed extraction")
        sources.append(metadata)
    hashes = {source["sha256"] for source in sources}
    rules = data.get("rateRules", [])
    for rule in rules:
        if rule.get("sourceSha256") not in hashes or not rule.get("sourceLocator"):
            raise ValueError("Every rule needs a source hash and row/page locator")
        if re.search(r"\b(other than|except|subject to|provided that)\b", rule.get("description", ""), re.I):
            rule["conditional"] = True
    if data.get("taxBasis") != "total_gst":
        raise ValueError("Rates must explicitly declare taxBasis=total_gst, not a CGST-only schedule")
    return rules, sources


def build_hsn(directory_path, output_path, *, code_column, description_column, directory_url,
              rates_path=None, retrieved_at=None):
    directory_path, output_path = Path(directory_path), Path(output_path)
    if output_path.exists():
        raise ValueError("Refusing to overwrite an existing versioned artifact")
    stamp = datetime.now(timezone.utc).isoformat()
    if retrieved_at is None:
        raise ValueError("Supply the actual source retrieved_at timestamp; build time is not retrieval time")
    rows = iter(directory_rows(directory_path))
    headers = next(rows, ())
    if headers.count(code_column) != 1 or headers.count(description_column) != 1:
        raise ValueError("Map unique code and description column names explicitly")
    code_index, description_index = headers.index(code_column), headers.index(description_column)
    codes = {}
    for row_number, row in enumerate(rows, 2):
        if not any(value not in ("", None) for value in row):
            continue
        if row_number > 100001:
            raise ValueError("Directory exceeds 100,000 rows")
        if len(row) <= max(code_index, description_index):
            raise ValueError(f"Incomplete directory row {row_number}")
        code, description = row[code_index], row[description_index]
        if not isinstance(code, str):
            raise ValueError(f"HSN code at row {row_number} must be text; preserve leading zeros explicitly")
        code = re.sub(r"\s+", "", code.strip())
        if not re.fullmatch(r"[0-9]{2,8}", code) or not isinstance(description, str) or not description.strip():
            raise ValueError(f"Invalid directory code/description at row {row_number}")
        entry = dict(code=code, description=description.strip(), chapter=code[:2])
        # Old-format rows embed spaces inside the code and a "~" separator in the
        # description; the same code also appears as a clean row. On conflict
        # prefer the description without "~".
        if code in codes and codes[code] != entry:
            current, incoming = codes[code]["description"], entry["description"]
            if "~" in current and "~" not in incoming:
                codes[code] = entry
            elif "~" in incoming and "~" not in current:
                pass
            else:
                raise ValueError(f"Conflicting descriptions for code {code}")
        else:
            codes[code] = entry
    if not codes:
        raise ValueError("Directory contains no codes")
    sources = [source_metadata(directory_path, "GST HSN directory", directory_url, retrieved_at)]
    rules = []
    if rates_path is not None:
        rules, rate_sources = reviewed_rates(Path(rates_path))
        sources.extend(rate_sources)
    artifact = HsnArtifact(
        version=datetime.now(timezone.utc).strftime("%Y%m%d"), builtAt=stamp,
        sources=sources, codes=sorted(codes.values(), key=lambda item: item["code"]),
        rateRules=rules,
    ).model_dump(exclude_none=True)
    compressed = gzip.compress(json.dumps(artifact, ensure_ascii=True, separators=(",", ":")).encode(), mtime=0)
    if len(compressed) >= MAX_ARTIFACT_BYTES:
        raise ValueError("Artifact must stay below 5 MiB; supply a documented chapter subset")
    with output_path.open("xb") as stream:
        stream.write(compressed)
    return artifact
